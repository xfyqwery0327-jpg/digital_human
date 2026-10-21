# -*- coding: utf-8 -*-
"""
build_rag.py - RAG 索引离线构建脚本（Windows 版）

用法:
    conda activate robot
    cd C:\\ro
    python build_rag.py
"""
import os
import re
import sys
import json
import pickle
from pathlib import Path
from collections import Counter

# ---------- Windows 控制台 UTF-8 兼容（防止 emoji/中文报错） ----------
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass

import faiss
import numpy as np
from sentence_transformers import SentenceTransformer
from docx import Document

# ==================== 配置（路径全部锚定到脚本所在目录，与运行位置无关） ====================
BASE_DIR   = os.path.dirname(os.path.abspath(__file__))          # C:\ro
MODEL_DIR  = os.path.join(BASE_DIR, "models")
EMBED_PATH = os.path.join(MODEL_DIR, "iic", "nlp_gte_sentence-embedding_chinese-base")

# 知识库目录：按顺序查找，用第一个包含 .docx 的目录
DOCS_DIR_CANDIDATES = [
    os.path.join(BASE_DIR, "knowledge_base"),   # C:\ro\knowledge_base（推荐位置）
    r"C:\knowledge_base",                        # C 盘根目录（你当前放文件的位置）
]

# 文件名关键词 → category 映射
CATEGORY_MAP = {
    "魔墙": "headquarters",
    "讲解词": "headquarters",
    "宽创": "headquarters",
    "妇好": "exhibition_fuhao",
    "B1": "exhibition_wanlijianghe",
    "万里江河": "exhibition_wanlijianghe",
}

# ==================== 工具函数 ====================
def resolve_docs_dir():
    """按顺序查找知识库目录，返回第一个包含 .docx 的目录"""
    for d in DOCS_DIR_CANDIDATES:
        if os.path.isdir(d):
            files = [f for f in Path(d).rglob("*.docx") if not f.name.startswith("~$")]
            if files:
                return d
    return None

def read_docx_text(path):
    """读取 docx 的全部文字（含表格内容）"""
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # 表格里的文字也提取（讲解词里常有表格）
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)

def load_docs(docs_dir):
    """扫描目录下所有 .docx 文件"""
    docs = []
    for f in Path(docs_dir).rglob("*.docx"):
        if f.name.startswith("~$"):   # 跳过 Word 临时文件
            continue
        try:
            text = read_docx_text(f)
            if not text.strip():
                print(f"⚠️  跳过空文档: {f.name}")
                continue

            # 根据文件名匹配分类
            category = "general"
            for keyword, cat in CATEGORY_MAP.items():
                if keyword in f.stem:
                    category = cat
                    break

            docs.append({"text": text, "source": f.name, "category": category})
            print(f"✅ 加载: {f.name} ({len(text)} 字) → {category}")
        except Exception as e:
            print(f"❌ 读取 {f.name} 失败: {e}")

    # 提醒：.doc 老格式无法读取
    for f in Path(docs_dir).rglob("*.doc"):
        if f.suffix.lower() == ".doc":
            print(f"⚠️  发现老格式 .doc 文件，无法读取，请用 Word 另存为 .docx: {f.name}")
    return docs

def is_question(line):
    """判断一行是否为问题"""
    return line.endswith("？") or line.endswith("?") or (line.startswith("你") and len(line) <= 30)

def extract_qa_pairs(text):
    """
    从问答库中提取 Q&A 对。
    docx 读出的文本是单换行分隔，因此按行扫描：
    问题行（以？结尾或"你"开头的短句）+ 后续若干行 = 一个问答对
    """
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    qa_pairs = []
    i = 0
    while i < len(lines):
        if is_question(lines[i]):
            question = lines[i]
            i += 1
            answer_parts = []
            while i < len(lines) and not is_question(lines[i]):
                answer_parts.append(lines[i])
                i += 1
            if answer_parts:
                qa_pairs.append((question, "\n".join(answer_parts)))
        else:
            i += 1
    return qa_pairs

def detect_venue(section_title, source):
    """根据章节标题自动识别展区"""
    # 总部展厅展区
    if "荣誉墙" in section_title or "铜牌墙" in section_title:
        return "honor_wall"
    elif "魔墙" in section_title or "发展历程" in section_title:
        return "magic_wall"
    elif "设计院" in section_title:
        return "design_institute"
    elif "负一楼" in section_title or "人工智能" in section_title:
        return "ai_tech_hall"
    # 妇好展厅主题
    elif "军事" in section_title or "征战" in section_title:
        return "military"
    elif "祭祀" in section_title or "神职" in section_title:
        return "ritual"
    elif "文物" in section_title or "妇好墓" in section_title:
        return "artifact"
    elif "问答" in section_title or "你是谁" in section_title:
        return "qa_library"
    # B1展厅朝代
    elif "先秦" in section_title or "远古" in section_title:
        return "pre_qin"
    elif "秦汉" in section_title or "兵马俑" in section_title:
        return "qin_han"
    elif "魏晋" in section_title or "南北朝" in section_title:
        return "wei_jin"
    elif "隋唐" in section_title or "唐代" in section_title:
        return "sui_tang"
    elif "宋元明清" in section_title or "清明上河图" in section_title:
        return "song_yuan_ming_qing"
    elif "AI" in section_title or "智能体" in section_title or "科技" in section_title:
        return "ai_interactive"
    return "general"

def chunk_text_fixed(text, size=500, overlap=100):
    """固定长度切块，带重叠"""
    result, start = [], 0
    while start < len(text):
        result.append(text[start:start + size])
        start += size - overlap
    return result

def chunk_by_section(text, source, category):
    """
    智能章节拆分：
    1. 特殊处理妇好知识库的问答部分（整对保留，不切碎）
    2. 按标题层级拆分（一、二、三... / 1. 2. 3... / #### / &&&&）
    3. 小章节（≤600字）整块保留，大章节切块（500字/块，重叠100字）
    """
    chunks = []

    # ========== 特殊处理：妇好知识库的问答库 ==========
    if "妇好" in source and "你是谁？" in text:
        qa_start = text.find("你是谁？")
        if qa_start > 0:
            main_content = text[:qa_start]
            qa_content = text[qa_start:]

            qa_pairs = extract_qa_pairs(qa_content)
            print(f"  └─ 提取了 {len(qa_pairs)} 个问答对")

            for q, a in qa_pairs:
                chunks.append({
                    "text": f"问：{q}\n答：{a}",
                    "source": source,
                    "category": category,
                    "section": "AI数字人问答库",
                    "venue": "qa_library",
                    "type": "qa",
                })
            text = main_content  # 剩余主体内容继续走章节拆分

    # ========== 按标题层级拆分 ==========
    section_pattern = r'''
        (?:
            \n(?=[一二三四五六七八九十百千]+、)          # 一、二、三、
            |\n(?=\d+[.、）)])                           # 1. 2） 3、
            |\n(?=\#+\s)                                 # ### 标题
            |\n(?=&&&&)                                  # &&&&
            |\n(?=第[一二三四五六七八九十]+[页部章节])    # 第一页
            |\n(?=【)                                    # 【标题】
        )
    '''
    sections = re.split(section_pattern, text, flags=re.VERBOSE)

    for sec in sections:
        sec = sec.strip()
        if len(sec) < 50:   # 跳过太短的片段
            continue

        # 提取章节标题（第一行前40字）
        first_line = sec.split("\n", 1)[0][:40].strip()
        section_title = first_line if first_line else sec[:30] + "..."

        venue = detect_venue(section_title, source)

        if len(sec) <= 600:
            # 小章节整块保留
            chunks.append({
                "text": sec,
                "source": source,
                "category": category,
                "section": section_title,
                "venue": venue,
                "type": "section",
            })
        else:
            # 大章节切块
            for idx, sub in enumerate(chunk_text_fixed(sec, size=500, overlap=100)):
                chunks.append({
                    "text": sub,
                    "source": source,
                    "category": category,
                    "section": f"{section_title} (第{idx+1}段)",
                    "venue": venue,
                    "type": "chunk",
                })
    return chunks

def resolve_embed_path():
    """定位 embedding 模型目录（兼容 modelscope 不同的缓存结构）"""
    if os.path.exists(EMBED_PATH):
        return EMBED_PATH
    if os.path.isdir(MODEL_DIR):
        for root, dirs, files in os.walk(MODEL_DIR):
            if "config.json" in files and "embedding" in root.lower().replace("\\", "/"):
                return root
    return None

def encode_texts(embedder, texts):
    """编码文本并归一化"""
    embs = embedder.encode(texts, show_progress_bar=True, batch_size=32).astype("float32")
    faiss.normalize_L2(embs)
    return embs

def print_stats(chunks):
    """打印统计信息"""
    print("\n" + "=" * 60)
    print("📊 索引统计")
    print("=" * 60)
    print(f"总块数: {len(chunks)}")

    print("\n按类别分布:")
    for cat, count in Counter(c["category"] for c in chunks).most_common():
        print(f"  • {cat}: {count} 块")

    print("\n按类型分布:")
    for t, count in Counter(c["type"] for c in chunks).most_common():
        print(f"  • {t}: {count} 块")

    venue_stats = Counter(c["venue"] for c in chunks if c["venue"] != "general")
    if venue_stats:
        print("\n按展区/主题分布:")
        for v, count in venue_stats.most_common():
            print(f"  • {v}: {count} 块")

    print("\n按文档分布:")
    for src, count in Counter(c["source"] for c in chunks).most_common():
        print(f"  • {src}: {count} 块")
    print("=" * 60 + "\n")

# ==================== 主流程 ====================
def main():
    print("\n" + "=" * 60)
    print("🏗️  开始构建 RAG 索引（Windows 优化版）")
    print("=" * 60 + "\n")

    # 1. 定位知识库目录
    docs_dir = resolve_docs_dir()
    if docs_dir is None:
        default_dir = DOCS_DIR_CANDIDATES[0]
        os.makedirs(default_dir, exist_ok=True)
        print("❌ 没有找到包含 .docx 文件的知识库目录。")
        print("请把 .docx 文件放入以下任一目录后重新运行：")
        for d in DOCS_DIR_CANDIDATES:
            print(f"   • {d}")
        return
    print(f"📂 使用知识库目录: {docs_dir}\n")

    # 2. 加载文档
    print("📖 扫描文档...")
    docs = load_docs(docs_dir)
    if not docs:
        print(f"\n❌ {docs_dir} 中没有可读取的 .docx 文件。")
        return
    print(f"\n✅ 加载了 {len(docs)} 个文档\n")

    # 3. 拆分成块
    print("✂️  智能拆分文本块...")
    chunks = []
    for doc in docs:
        doc_chunks = chunk_by_section(doc["text"], doc["source"], doc["category"])
        chunks.extend(doc_chunks)
        print(f"  • {doc['source']}: {len(doc_chunks)} 块")

    if not chunks:
        print("\n❌ 没有有效的文本块，请检查文档内容。")
        return

    # 4. 加载 embedding 模型
    print("\n🧠 加载 embedding 模型...")
    embed_path = resolve_embed_path()
    if embed_path is None:
        print(f"❌ 模型不存在: {EMBED_PATH}")
        print("请先下载模型：")
        print(f'  python -c "from modelscope import snapshot_download; '
              f"snapshot_download('iic/nlp_gte_sentence-embedding_chinese-base', "
              f"cache_dir=r'{MODEL_DIR}')\"")
        return
    print(f"   模型路径: {embed_path}")
    embedder = SentenceTransformer(embed_path)

    # 5. 编码文本
    print("\n🔢 编码文本向量...")
    texts = [c["text"] for c in chunks]
    embs = encode_texts(embedder, texts)

    # 6. 构建 faiss 索引
    print("\n🔍 构建 faiss 索引...")
    dim = embs.shape[1]
    index = faiss.IndexFlatIP(dim)   # 内积检索（归一化后=余弦相似度）
    index.add(embs)

    # 7. 保存索引（修改为 faiss + JSON 格式）
    OUTPUT_DIR = Path(BASE_DIR) / "data" / "faiss_index"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f"\n💾 保存索引到: {OUTPUT_DIR}")

    # 保存 faiss 索引
    faiss.write_index(index, str(OUTPUT_DIR / "index.faiss"))

    # 保存文档块（JSON 格式）
    with open(OUTPUT_DIR / "chunks.json", "w", encoding="utf-8") as f:
        json.dump(chunks, f, ensure_ascii=False, indent=2)

    INDEX_SIZE = (OUTPUT_DIR / "index.faiss").stat().st_size / 1024 / 1024
    CHUNKS_SIZE = (OUTPUT_DIR / "chunks.json").stat().st_size / 1024 / 1024

    # 8. 打印统计
    print_stats(chunks)

    print("✅ 索引构建完成！")
    print(f"索引文件: {OUTPUT_DIR / 'index.faiss'} ({INDEX_SIZE:.2f} MB)")
    print(f"文档块: {OUTPUT_DIR / 'chunks.json'} ({CHUNKS_SIZE:.2f} MB)\n")

if __name__ == "__main__":
    main()